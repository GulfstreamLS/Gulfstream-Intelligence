import io
import uuid
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from pydantic import BaseModel
from sqlalchemy import delete, func, select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1._audit import get_ip, log_audit
from app.db.session import get_db
from app.middleware.auth import check_active_subscription, get_current_user, get_user_or_none
from app.models.chat import Conversation, Message, MessageRole
from app.models.notification import Notification, NotificationType
from app.models.organization import MemberRole, OrganizationMember
from app.models.project import Project
from app.models.regulatory import AnalysisDocument, Gap
from app.models.simulation import SimulationSession
from app.models.user import User
from app.schemas.chat import ConversationResponse
from app.schemas.project import ProjectCreate, ProjectListResponse, ProjectResponse, ProjectUpdate

ALLOWED_DOC_EXTENSIONS = {"pdf", "docx", "doc", "txt"}

router = APIRouter(prefix="/projects", tags=["projects"])


def _org_id(user: Optional[User]) -> Optional[uuid.UUID]:
    return user.organization_id if user else None


async def _can_delete(resource_user_id: uuid.UUID, current_user: User, db: AsyncSession) -> bool:
    """Creator or org owner can delete."""
    if resource_user_id == current_user.id:
        return True
    if current_user.organization_id:
        result = await db.execute(
            select(OrganizationMember).where(
                OrganizationMember.user_id == current_user.id,
                OrganizationMember.organization_id == current_user.organization_id,
            )
        )
        member = result.scalar_one_or_none()
        return member is not None and member.role == MemberRole.OWNER
    return False


def _build_list_query(user: Optional[User]):
    if user and user.organization_id:
        return select(Project).where(Project.organization_id == user.organization_id)
    elif user:
        return select(Project).where(Project.user_id == user.id)
    else:
        return select(Project).where(Project.user_id == None)  # noqa: E711 — fallback


async def _notify_org_members(db: AsyncSession, org_id: uuid.UUID, exclude_id: uuid.UUID, notif_type: str, title: str, body: str, resource_id: str):
    from app.models.organization import MemberStatus
    result = await db.execute(
        select(OrganizationMember).where(
            OrganizationMember.organization_id == org_id,
            OrganizationMember.user_id != exclude_id,
            OrganizationMember.status == MemberStatus.ACTIVE,
        )
    )
    for m in result.scalars().all():
        db.add(Notification(
            user_id=m.user_id,
            type=notif_type,
            title=title,
            body=body,
            resource_type="project",
            resource_id=resource_id,
        ))


@router.get("", response_model=ProjectListResponse)
async def list_projects(
    page: int = 1,
    page_size: int = 10,
    search: Optional[str] = None,
    status_filter: Optional[str] = None,
    authority: Optional[str] = None,
    current_user: Optional[User] = Depends(get_user_or_none),
    db: AsyncSession = Depends(get_db),
):
    q = _build_list_query(current_user)

    if search:
        term = f"%{search}%"
        q = q.where(
            Project.name.ilike(term)
            | Project.indication.ilike(term)
            | Project.therapeutic_area.ilike(term)
        )
    if status_filter:
        q = q.where(Project.status == status_filter)
    if authority:
        q = q.where(Project.authorities.contains([authority]))

    count_q = select(func.count()).select_from(q.subquery())
    total = (await db.execute(count_q)).scalar_one()
    q = q.order_by(Project.updated_at.desc()).offset((page - 1) * page_size).limit(page_size)
    rows = (await db.execute(q)).scalars().all()
    return ProjectListResponse(items=list(rows), total=total, page=page, page_size=page_size)


@router.post("", response_model=ProjectResponse, status_code=status.HTTP_201_CREATED)
async def create_project(
    request: Request,
    data: ProjectCreate,
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    project = Project(
        user_id=current_user.id,
        organization_id=_org_id(current_user),
        **data.model_dump(),
    )
    db.add(project)
    await db.commit()
    await db.refresh(project)
    await log_audit(db, current_user.id, "PROJECT_CREATED", resource_type="project", resource_id=project.id, resource_name=project.name, ip_address=get_ip(request), organization_id=_org_id(current_user))

    if current_user.organization_id:
        await _notify_org_members(db, current_user.organization_id, current_user.id, NotificationType.PROJECT_CREATED, f"New project: {project.name}", f"{current_user.full_name or current_user.email} created project \"{project.name}\"", str(project.id))

    await db.commit()
    return project


@router.get("/{project_id}", response_model=ProjectResponse)
async def get_project(
    project_id: uuid.UUID,
    current_user: Optional[User] = Depends(get_user_or_none),
    db: AsyncSession = Depends(get_db),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    # Access check: own project OR same org
    if current_user:
        if current_user.organization_id and project.organization_id == current_user.organization_id:
            return project
        if project.user_id == current_user.id:
            return project
    raise HTTPException(status_code=404, detail="Project not found")


@router.patch("/{project_id}", response_model=ProjectResponse)
async def update_project(
    request: Request,
    project_id: uuid.UUID,
    data: ProjectUpdate,
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    # Access check
    accessible = (project.user_id == current_user.id) or (current_user.organization_id and project.organization_id == current_user.organization_id)
    if not accessible:
        raise HTTPException(status_code=404, detail="Project not found")

    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(project, field, value)
    await db.commit()
    await db.refresh(project)
    await log_audit(db, current_user.id, "PROJECT_UPDATED", resource_type="project", resource_id=project.id, resource_name=project.name, ip_address=get_ip(request), organization_id=_org_id(current_user))

    if current_user.organization_id:
        await _notify_org_members(db, current_user.organization_id, current_user.id, NotificationType.PROJECT_UPDATED, f"Project updated: {project.name}", f"{current_user.full_name or current_user.email} updated project \"{project.name}\"", str(project.id))

    await db.commit()
    return project


@router.delete("/{project_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_project(
    request: Request,
    project_id: uuid.UUID,
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Access check: must be in same org or own it
    accessible = (project.user_id == current_user.id) or (current_user.organization_id and project.organization_id == current_user.organization_id)
    if not accessible:
        raise HTTPException(status_code=404, detail="Project not found")

    # Delete permission: creator or org owner only
    if not await _can_delete(project.user_id, current_user, db):
        raise HTTPException(status_code=403, detail="Only the project creator or organization owner can delete this project")

    project_name = project.name
    await db.execute(delete(SimulationSession).where(SimulationSession.project_id == project_id))
    await db.execute(delete(Conversation).where(Conversation.project_id == project_id))
    await db.delete(project)

    await log_audit(db, current_user.id, "PROJECT_DELETED", resource_type="project", resource_id=project_id, resource_name=project_name, ip_address=get_ip(request), organization_id=_org_id(current_user))

    if current_user.organization_id:
        await _notify_org_members(db, current_user.organization_id, current_user.id, NotificationType.PROJECT_DELETED, f"Project deleted: {project_name}", f"{current_user.full_name or current_user.email} deleted project \"{project_name}\"", str(project_id))

    await db.commit()


@router.get("/{project_id}/conversations", response_model=list[ConversationResponse])
async def get_project_conversations(
    project_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    accessible = (project.user_id == current_user.id) or (current_user.organization_id and project.organization_id == current_user.organization_id)
    if not accessible:
        raise HTTPException(status_code=404, detail="Project not found")

    result = await db.execute(
        select(Conversation)
        .where(
            Conversation.project_id == project_id,
            Conversation.is_temporary == False,  # noqa: E712
        )
        .options(selectinload(Conversation.messages), selectinload(Conversation.project), selectinload(Conversation.user))
        .order_by(Conversation.updated_at.desc())
    )
    return result.scalars().all()


class _BulkDeleteBody(BaseModel):
    ids: list[uuid.UUID]


@router.post("/bulk-delete")
async def bulk_delete_projects(
    request: Request,
    data: _BulkDeleteBody,
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    """Delete multiple projects in a single request."""
    if not data.ids:
        return {"deleted": 0, "skipped": 0}

    org_id = _org_id(current_user)

    is_org_owner = False
    if org_id:
        result = await db.execute(
            select(OrganizationMember).where(
                OrganizationMember.user_id == current_user.id,
                OrganizationMember.organization_id == org_id,
            )
        )
        member = result.scalar_one_or_none()
        is_org_owner = member is not None and member.role == MemberRole.OWNER

    q = _build_list_query(current_user).where(Project.id.in_(data.ids))
    result = await db.execute(q)
    projects = result.scalars().all()

    deleted = skipped = 0
    for project in projects:
        can_del = project.user_id == current_user.id or is_org_owner
        if can_del:
            await db.execute(delete(Conversation).where(Conversation.project_id == project.id))
            await db.delete(project)
            deleted += 1
        else:
            skipped += 1

    if deleted:
        await log_audit(
            db, current_user.id, "PROJECT_BULK_DELETED",
            resource_type="project",
            resource_name=f"{deleted} projects",
            ip_address=get_ip(request),
            organization_id=org_id,
        )

    await db.commit()
    return {"deleted": deleted, "skipped": skipped}


@router.post("/import", response_model=dict)
async def import_projects(
    file: UploadFile = File(...),
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="File is empty")

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]

    def col(row, name):
        try:
            idx = headers.index(name)
            return row[idx]
        except ValueError:
            return None

    created, errors = 0, []
    for i, row in enumerate(rows[1:], start=2):
        name = col(row, "name")
        if not name:
            errors.append(f"Row {i}: missing Name")
            continue
        raw_auth = col(row, "authorities")
        authorities = [a.strip() for a in str(raw_auth).split(",")] if raw_auth else []
        raw_score = col(row, "readiness score")
        try:
            readiness_score = int(raw_score) if raw_score is not None else 0
        except (ValueError, TypeError):
            readiness_score = 0

        project = Project(
            user_id=current_user.id,
            organization_id=_org_id(current_user),
            name=str(name),
            type=str(col(row, "type") or "IND"),
            indication=str(col(row, "indication") or "") or None,
            therapeutic_area=str(col(row, "therapeutic area") or "") or None,
            dev_phase=str(col(row, "dev phase") or "") or None,
            status=str(col(row, "status") or "Planning"),
            readiness_score=readiness_score,
            authorities=authorities if authorities else None,
            product_type=str(col(row, "product type") or "") or None,
        )
        db.add(project)
        created += 1

    await db.commit()
    return {"created": created, "errors": errors}


# ── Project documents ─────────────────────────────────────────────────────────

class ProjectDocumentResponse(BaseModel):
    id: uuid.UUID
    project_id: Optional[uuid.UUID]
    conversation_id: Optional[uuid.UUID]
    filename: str
    file_type: str
    authority: Optional[str]
    summary: Optional[str]
    source: Optional[str] = None
    gap_count: int = 0
    created_at: datetime

    model_config = {"from_attributes": True}


def _project_document_response(doc: AnalysisDocument) -> ProjectDocumentResponse:
    if doc.conversation_id:
        source = "Regulatory Chat"
    elif doc.project_id and doc.gaps:
        source = "Gap Assessment"
    elif doc.project_id:
        source = "Project Upload"
    else:
        source = "Standalone Upload"
    return ProjectDocumentResponse(
        id=doc.id,
        project_id=doc.project_id,
        conversation_id=doc.conversation_id,
        filename=doc.filename,
        file_type=doc.file_type,
        authority=doc.authority,
        summary=doc.summary,
        source=source,
        gap_count=len(doc.gaps or []),
        created_at=doc.created_at,
    )


class ProjectSourceCountsResponse(BaseModel):
    prior_gap_assessment: int
    chat_outputs: int


async def _get_accessible_project(
    project_id: uuid.UUID, current_user: User, db: AsyncSession
) -> Project:
    project = await db.get(Project, project_id)
    accessible = project and (
        project.user_id == current_user.id
        or (current_user.organization_id and project.organization_id == current_user.organization_id)
    )
    if not accessible:
        raise HTTPException(status_code=404, detail="Project not found")
    return project


@router.post("/documents", response_model=ProjectDocumentResponse)
async def upload_standalone_document(
    request: Request,
    file: UploadFile = File(...),
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    """Upload an unlinked document for one-off standalone simulations."""
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED_DOC_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload PDF, Word, or TXT documents.",
        )

    content = await file.read()
    from app.services.document_processor import document_processor

    try:
        text = document_processor.extract_text(content, ext)
    except Exception:
        text = ""

    doc = AnalysisDocument(
        user_id=current_user.id,
        project_id=None,
        organization_id=_org_id(current_user),
        filename=file.filename,
        file_type=ext,
        file_path=f"uploads/{file.filename}",
        extracted_text=text or None,
        summary=(text[:500] if text else None),
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)

    await log_audit(
        db, current_user.id, "STANDALONE_DOCUMENT_UPLOADED",
        resource_type="document",
        resource_id=doc.id,
        resource_name=file.filename,
        ip_address=get_ip(request),
        organization_id=_org_id(current_user),
    )
    await db.commit()
    return doc


@router.get("/{project_id}/documents", response_model=list[ProjectDocumentResponse])
async def list_project_documents(
    project_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List documents saved to a project."""
    await _get_accessible_project(project_id, current_user, db)
    result = await db.execute(
        select(AnalysisDocument)
        .options(selectinload(AnalysisDocument.gaps))
        .where(AnalysisDocument.project_id == project_id)
        .order_by(AnalysisDocument.created_at.desc())
    )
    return [_project_document_response(doc) for doc in result.scalars().unique().all()]


@router.get("/{project_id}/source-counts", response_model=ProjectSourceCountsResponse)
async def get_project_source_counts(
    project_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Return available project source counts for HA simulation source controls."""
    await _get_accessible_project(project_id, current_user, db)

    gap_count = (await db.execute(
        select(func.count(Gap.id))
        .join(AnalysisDocument, Gap.document_id == AnalysisDocument.id)
        .where(AnalysisDocument.project_id == project_id)
    )).scalar_one()

    chat_output_count = (await db.execute(
        select(func.count(Message.id))
        .join(Conversation, Message.conversation_id == Conversation.id)
        .where(
            Conversation.project_id == project_id,
            Conversation.is_temporary == False,  # noqa: E712
            Message.role == MessageRole.ASSISTANT,
            Message.content != "",
        )
    )).scalar_one()

    return ProjectSourceCountsResponse(
        prior_gap_assessment=gap_count,
        chat_outputs=chat_output_count,
    )


@router.delete("/documents/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_standalone_document(
    request: Request,
    document_id: uuid.UUID,
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    """Delete an unlinked document used for one-off simulations."""
    doc = await db.get(AnalysisDocument, document_id)
    accessible = doc and doc.project_id is None and (
        doc.user_id == current_user.id
        or (current_user.organization_id and doc.organization_id == current_user.organization_id)
    )
    if not accessible:
        raise HTTPException(status_code=404, detail="Document not found")

    await db.delete(doc)
    await log_audit(
        db, current_user.id, "STANDALONE_DOCUMENT_DELETED",
        resource_type="document",
        resource_id=document_id,
        resource_name=doc.filename,
        ip_address=get_ip(request),
        organization_id=_org_id(current_user),
    )
    await db.commit()


@router.post("/{project_id}/documents", response_model=ProjectDocumentResponse)
async def upload_project_document(
    request: Request,
    project_id: uuid.UUID,
    file: UploadFile = File(...),
    save_to_project: bool = Form(True),
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    """Upload a document for a project.

    When ``save_to_project`` is true the document is linked to the project and
    becomes available to other features (Gap Assessment, Document Intelligence,
    Regulatory Chat). When false it is stored unlinked for one-off simulation use.
    """
    await _get_accessible_project(project_id, current_user, db)

    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED_DOC_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload PDF, Word, or TXT documents.",
        )

    content = await file.read()
    from app.services.document_processor import document_processor

    try:
        text = document_processor.extract_text(content, ext)
    except Exception:
        text = ""

    doc = AnalysisDocument(
        user_id=current_user.id,
        project_id=project_id if save_to_project else None,
        organization_id=_org_id(current_user),
        filename=file.filename,
        file_type=ext,
        file_path=f"uploads/{file.filename}",
        extracted_text=text or None,
        summary=(text[:500] if text else None),
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)

    await log_audit(
        db, current_user.id, "PROJECT_DOCUMENT_UPLOADED",
        resource_type="document",
        resource_id=doc.id,
        resource_name=file.filename,
        ip_address=get_ip(request),
        organization_id=_org_id(current_user),
    )
    await db.commit()
    return doc


@router.delete("/{project_id}/documents/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_project_document(
    request: Request,
    project_id: uuid.UUID,
    document_id: uuid.UUID,
    current_user: User = Depends(check_active_subscription),
    db: AsyncSession = Depends(get_db),
):
    """Delete a document saved to a project."""
    await _get_accessible_project(project_id, current_user, db)
    doc = await db.get(AnalysisDocument, document_id)
    if not doc or doc.project_id != project_id:
        raise HTTPException(status_code=404, detail="Document not found")

    await db.delete(doc)
    await log_audit(
        db, current_user.id, "PROJECT_DOCUMENT_DELETED",
        resource_type="document",
        resource_id=document_id,
        resource_name=doc.filename,
        ip_address=get_ip(request),
        organization_id=_org_id(current_user),
    )
    await db.commit()
